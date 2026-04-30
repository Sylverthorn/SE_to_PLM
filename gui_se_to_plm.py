import sys
import os
import time
import pythoncom
from datetime import datetime
from PyQt5.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
                             QLabel, QLineEdit, QPushButton, QTextEdit, QFileDialog, QMessageBox, QComboBox, QDialog, QProgressBar)
from PyQt5.QtCore import QThread, pyqtSignal, Qt
from PyQt5.QtGui import QTextCursor, QColor
from concurrent.futures import ThreadPoolExecutor

from se_to_plm import indexer_les_plans_projet_entier, extraire_metadonnees, extraire_metadonnees_rapide, determiner_classe
import win32com.client
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

class ExtractionThread(QThread):
    log_signal = pyqtSignal(str, str)
    finished_signal = pyqtSignal()
    progress_signal = pyqtSignal(int, int, str)  # (valeur, maximum, message)

    def __init__(self, chemin_asm, dossier_sortie, nom_sortie, dossier_dft=None, mode_recherche_dft="les_deux"):
        super().__init__()
        self.chemin_asm = chemin_asm
        self.dossier_sortie = dossier_sortie
        self.nom_sortie = nom_sortie
        self.dossier_dft = dossier_dft
        self.mode_recherche_dft = mode_recherche_dft
        self._cancelled = False
    
    def cancel(self):
        self._cancelled = True

    def run(self):
        try:
            self.log_signal.emit("=" * 60, 'info')
            self.log_signal.emit("Début de l'extraction PLM", 'info')
            self.log_signal.emit("=" * 60, 'info')
            self.progress_signal.emit(0, 100, "Initialisation...")

            # Callback pour la progression de l'indexation
            def on_index_progress(scanned, found, total):
                if total > 0:
                    pct = min(10, int((scanned / total) * 10))  # Max 10% pour l'indexation
                    self.progress_signal.emit(pct, 100, f"Indexation: {scanned} dossiers scannés, {found} plans trouvés")

            self.log_signal.emit(f"\n--- Indexation des plans (.dft) [Mode: {self.mode_recherche_dft}] ---", 'info')
            self.progress_signal.emit(0, 100, "Indexation des plans...")
            index_plans = indexer_les_plans_projet_entier(self.chemin_asm, self.dossier_dft, self.mode_recherche_dft, callback_progress=on_index_progress)
            self.log_signal.emit(f"-> {len(index_plans)} plan(s) détecté(s).", 'info')
            self.progress_signal.emit(10, 100, f"{len(index_plans)} plans indexés")
            
            self.log_signal.emit("\nConnexion à Solid Edge...", 'info')
            try:
                # 1. Tente de se brancher sur un Solid Edge déjà ouvert (Instantané)
                app = win32com.client.GetActiveObject("SolidEdge.Application")
                self.log_signal.emit("Connecté à l'instance existante de Solid Edge.", 'success')
            except pythoncom.com_error:
                # 2. S'il n'est pas ouvert, on le lance (Prend quelques secondes)
                self.log_signal.emit("Démarrage de Solid Edge en arrière-plan...", 'info')
                app = win32com.client.dynamic.Dispatch("SolidEdge.Application")
                app.Visible = False
                self.log_signal.emit("Solid Edge démarré.", 'success')
            
            # Désactiver les alertes pour accélérer l'ouverture des fichiers
            app.DisplayAlerts = False
            
            doc_racine = app.Documents.Open(self.chemin_asm)
            # COM API bloque jusqu'à ce que le document soit chargé, pas besoin de sleep
            self.log_signal.emit("Document chargé.", 'success')
            
            lignes_excel = []
            compteur_ordre = 1
            liste_plans_a_rajouter = []
            stats = {"3d": 0, "2d": 0}
            
            def get_suffixe_fichier(nom_fichier):
                """Retourne le suffixe approprié selon l'extension du fichier."""
                ext = os.path.splitext(nom_fichier)[1].lower()
                if ext == '.asm':
                    return "(ASM)"
                elif ext in ['.par', '.psm']:
                    return "(PRT)"
                elif ext == '.dft':
                    return "(DRW)"
                return ""
            
            def ajouter_ligne(niveau, relation, nom_fichier, chemin_complet, classe, qte=1, rev="1", desig="", ver="-"):
                nonlocal compteur_ordre
                ref_util = os.path.splitext(nom_fichier)[0]
                special_cad = os.path.splitext(nom_fichier)[0]  # Sans extension
                suffixe = get_suffixe_fichier(nom_fichier)
                # Normaliser le chemin (remplace / par \) et concaténer sans espace
                chemin_normalise = os.path.normpath(chemin_complet)
                attachement = f"{chemin_normalise}{suffixe}" if suffixe else chemin_normalise

                lignes_excel.append([
                    niveau, relation, compteur_ordre, qte, "", special_cad,
                    classe, ref_util, ver, rev, desig, "", attachement
                ])
                compteur_ordre += 1
            
            def explorer_occurrences(occurrences, niveau):
                nonlocal stats
                if occurrences is None: return
                
                dict_occ = {}
                for i in range(1, occurrences.Count + 1):
                    try:
                        occ = occurrences.Item(i)
                        path_reel = ""
                        nom_reel = ""
                        try:
                            path_reel = occ.OccurrenceDocument.FullName
                            nom_reel = os.path.basename(path_reel)
                        except:
                            nom_reel = occ.Name.split(':')[0]
                        
                        if nom_reel not in dict_occ:
                            dict_occ[nom_reel] = {"qte": 1, "obj": occ, "chemin": path_reel}
                        else:
                            dict_occ[nom_reel]["qte"] += 1
                    except: continue
                
                for nom, data in dict_occ.items():
                    stats["3d"] += 1
                    classe_3d = determiner_classe(nom)
                    # Utiliser la méthode rapide avec le chemin du fichier
                    meta = extraire_metadonnees_rapide(data["chemin"])
                    
                    # On ajoute uniquement le 3D dans l'arbre principal
                    ajouter_ligne(niveau, "ComposedOf", nom, data["chemin"], classe_3d, data["qte"], meta["revision"], meta["designation"], meta["version"])
                    
                    nom_sans_ext = os.path.splitext(nom)[0].lower()
                    if nom_sans_ext in index_plans:
                        chemin_dft = index_plans[nom_sans_ext]
                        nom_dft = os.path.basename(chemin_dft)
                        
                        # On stocke pour l'ajouter à la fin du fichier
                        liste_plans_a_rajouter.append({
                            "dft_nom": nom_dft, "dft_path": chemin_dft,
                            "src_nom": nom, "src_path": data["chemin"],
                            "src_classe": classe_3d, "src_rev": meta["revision"], "src_desig": meta["designation"],
                            "src_ver": meta["version"]
                        })
                        stats["2d"] += 1
                    
                    if data["obj"].Subassembly:
                        try: explorer_occurrences(data["obj"].OccurrenceDocument.Occurrences, niveau + 1)
                        except: pass
            
            self.log_signal.emit("\nAnalyse de la structure...", 'info')
            
            # Utiliser la méthode rapide pour le document racine aussi
            meta_root = extraire_metadonnees_rapide(doc_racine.FullName)
            nom_root = os.path.basename(doc_racine.FullName)
            ajouter_ligne(0, "", nom_root, doc_racine.FullName, "SUB_ASSY_A", 1, meta_root["revision"], meta_root["designation"], meta_root["version"])
            
            nom_root_pur = os.path.splitext(nom_root)[0].lower()
            if nom_root_pur in index_plans:
                path_dft_root = index_plans[nom_root_pur]
                nom_dft_root = os.path.basename(path_dft_root)
                liste_plans_a_rajouter.append({
                    "dft_nom": nom_dft_root, "dft_path": path_dft_root,
                    "src_nom": nom_root, "src_path": doc_racine.FullName,
                    "src_classe": "SUB_ASSY_A", "src_rev": meta_root["revision"], "src_desig": meta_root["designation"],
                    "src_ver": meta_root["version"]
                })
                stats["2d"] += 1
            
            explorer_occurrences(doc_racine.Occurrences, 1)
            
            # --- A LA FIN : AJOUT DES PLANS EN LEVEL 0 ET LEURS 3D EN LEVEL 1 ---
            self.log_signal.emit("\nExtraction métadonnées des plans...", 'info')
            self.progress_signal.emit(50, 100, "Extraction des métadonnées des plans...")
            
            # Collecter les chemins uniques des plans à traiter
            plans_uniques = {}
            for item in liste_plans_a_rajouter:
                if item["dft_path"] not in plans_uniques:
                    plans_uniques[item["dft_path"]] = item
            
            plans_list = list(plans_uniques.values())
            total_plans = len(plans_list)
            plans_deja_traites = set()
            
            for idx, item in enumerate(plans_list):
                if self._cancelled:
                    break
                    
                # Mettre à jour la progression tous les 5 plans
                if idx % 5 == 0 and total_plans > 0:
                    progress_pct = 50 + int((idx / total_plans) * 20)  # 50-70%
                    self.progress_signal.emit(progress_pct, 100, f"Plan {idx+1}/{total_plans}: {item['dft_nom'][:30]}...")
                
                if item["dft_path"] not in plans_deja_traites:
                    # Utiliser la méthode ultra-rapide (FileProperties) au lieu d'ouvrir le document
                    meta_dft = extraire_metadonnees_rapide(item["dft_path"], debug=(len(plans_deja_traites) == 0))
                    
                    # Utiliser la designation de la pièce 3D associée pour le DFT
                    ajouter_ligne(0, "", item["dft_nom"], item["dft_path"], "CAD_DRAWING_A", 1, meta_dft["revision"], item["src_desig"], meta_dft["version"])
                    ajouter_ligne(1, "Drawing", item["src_nom"], item["src_path"], item["src_classe"], 1, item["src_rev"], item["src_desig"], item["src_ver"])
                    plans_deja_traites.add(item["dft_path"])
            
            self.log_signal.emit(f"Analyse terminée : {stats['3d']} fichiers 3D, {stats['2d']} plans", 'success')
            self.progress_signal.emit(70, 100, f"Analyse terminée: {stats['3d']} fichiers 3D, {stats['2d']} plans")
            
            self.log_signal.emit("\nGénération du fichier Excel...", 'info')
            self.progress_signal.emit(75, 100, "Génération du fichier Excel...")
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Structure"
            
            headers = ["Level", "Relationship", "ordre", "quantite", "repere", "SpecialCAD", "Class", "ref_utilisat", "version", "revision", "designation", "dia_se", "Attachments"]
            ws.append(headers)
            
            header_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
            orange_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
            
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill if cell.column < 13 else orange_fill
                cell.alignment = Alignment(horizontal="left")
            
            self.progress_signal.emit(80, 100, "Écriture des données...")
            for l in lignes_excel: 
                ws.append(l)
            
            # Optimisation: calcul des largeurs de colonnes en parallèle
            self.progress_signal.emit(90, 100, "Calcul des largeurs de colonnes...")
            columns = list(ws.columns)
            
            def calc_column_width(col_data):
                col_cells, idx = col_data
                max_length = 0
                for cell in col_cells:
                    try: 
                        max_length = max(max_length, len(str(cell.value)))
                    except: 
                        pass
                return (idx, max_length + 2)
            
            # Utiliser ThreadPoolExecutor pour paralléliser le calcul
            with ThreadPoolExecutor(max_workers=4) as executor:
                futures = [executor.submit(calc_column_width, (col, i)) for i, col in enumerate(columns)]
                for future in futures:
                    idx, width = future.result()
                    ws.column_dimensions[columns[idx][0].column_letter].width = width
            
            nom_sortie = self.nom_sortie
            if not nom_sortie.endswith('.xlsx'):
                nom_sortie += '.xlsx'
            
            chemin_complet = os.path.join(self.dossier_sortie, nom_sortie)
            self.progress_signal.emit(95, 100, "Sauvegarde du fichier...")
            wb.save(chemin_complet)
            
            self.progress_signal.emit(100, 100, "Terminé!")
            self.log_signal.emit(f"\nFichier généré : {chemin_complet}", 'success')
            self.log_signal.emit("=" * 60, 'info')
            self.log_signal.emit("Extraction terminée avec succès !", 'success')
            
        except Exception as e:
            self.log_signal.emit(f"\nErreur : {e}", 'error')
            self.log_signal.emit("=" * 60, 'error')
        
        finally:
            # On NE ferme PLUS Solid Edge ici pour pouvoir le réutiliser à la prochaine extraction !
            # Solid Edge sera fermé proprement quand l'utilisateur ferme l'application (closeEvent)
            self.finished_signal.emit()

class PLMExtractorGUI(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Extracteur PLM - Solid Edge")
        self.setGeometry(100, 100, 800, 600)
        self.extraction_en_cours = False
        
        self.appliquer_style()
        self.creer_interface()
    
    def closeEvent(self, event):
        """Fermer Solid Edge quand l'application est fermée avec une fenêtre de chargement."""
        # Créer une fenêtre de chargement
        loading_dialog = QDialog(self)
        loading_dialog.setWindowTitle("Fermeture")
        loading_dialog.setFixedSize(300, 100)
        layout = QVBoxLayout(loading_dialog)
        label = QLabel("Fermeture de Solid Edge...")
        label.setAlignment(Qt.AlignCenter)
        layout.addWidget(label)
        loading_dialog.show()
        
        # Forcer la mise à jour de l'interface
        QApplication.processEvents()
        
        try:
            app = win32com.client.dynamic.Dispatch("SolidEdge.Application")
            app.Quit()
        except:
            pass
        
        loading_dialog.close()
        event.accept()
    
    def appliquer_style(self):
        chemin_style = os.path.join(os.path.dirname(__file__), 'style.qss')
        if os.path.exists(chemin_style):
            with open(chemin_style, 'r', encoding='utf-8') as f:
                self.setStyleSheet(f.read())
        
    def creer_interface(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setSpacing(10)
        
        asm_layout = QHBoxLayout()
        asm_layout.addWidget(QLabel("Fichier ASM :"))
        self.chemin_asm_edit = QLineEdit()
        self.chemin_asm_edit.setReadOnly(True)
        asm_layout.addWidget(self.chemin_asm_edit)
        btn_parcourir = QPushButton("Parcourir...")
        btn_parcourir.clicked.connect(self.choisir_fichier_asm)
        asm_layout.addWidget(btn_parcourir)
        main_layout.addLayout(asm_layout)
        
        # Mode de recherche DFT
        mode_layout = QHBoxLayout()
        mode_layout.addWidget(QLabel("Mode recherche DFT :"))
        self.mode_dft_combo = QComboBox()
        self.mode_dft_combo.addItems([
            "Arborescence uniquement",
            "Dossier spécifique uniquement",
            "Les deux (arborescence + dossier)"
        ])
        self.mode_dft_combo.setCurrentIndex(2)  # "Les deux" par défaut
        self.mode_dft_combo.currentIndexChanged.connect(self.on_mode_dft_changed)
        mode_layout.addWidget(self.mode_dft_combo)
        main_layout.addLayout(mode_layout)

        # Dossier DFT spécifique
        dft_layout = QHBoxLayout()
        dft_layout.addWidget(QLabel("Dossier plans (.dft) :"))
        self.dossier_dft_edit = QLineEdit()
        self.dossier_dft_edit.setReadOnly(True)
        self.dossier_dft_edit.setPlaceholderText("Sélectionner un dossier...")
        self.dossier_dft_edit.setEnabled(False)  # Désactivé par défaut
        dft_layout.addWidget(self.dossier_dft_edit)
        self.btn_parcourir_dft = QPushButton("Parcourir...")
        self.btn_parcourir_dft.clicked.connect(self.choisir_dossier_dft)
        self.btn_parcourir_dft.setEnabled(False)  # Désactivé par défaut
        dft_layout.addWidget(self.btn_parcourir_dft)
        main_layout.addLayout(dft_layout)
        
        sortie_layout = QHBoxLayout()
        sortie_layout.addWidget(QLabel("Nom de sortie :"))
        self.nom_sortie_edit = QLineEdit(f"Export_PLM_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        sortie_layout.addWidget(self.nom_sortie_edit)
        main_layout.addLayout(sortie_layout)
        
        self.dossier_sortie = os.path.join(os.path.expanduser("~"), "Documents", "Exports_PLM")
        lbl_dossier = QLabel(f"Dossier : {self.dossier_sortie}")
        lbl_dossier.setStyleSheet("color: gray; font-style: italic;")
        main_layout.addWidget(lbl_dossier)
        
        self.btn_extraire = QPushButton("Lancer l'extraction")
        self.btn_extraire.clicked.connect(self.lancer_extraction)
        self.btn_extraire.setMinimumHeight(40)
        main_layout.addWidget(self.btn_extraire)
        
        # Barre de progression
        self.progress_bar = QProgressBar()
        self.progress_bar.setMinimum(0)
        self.progress_bar.setMaximum(100)
        self.progress_bar.setValue(0)
        self.progress_bar.setTextVisible(True)
        self.progress_bar.setFormat("%p% - %v")
        main_layout.addWidget(self.progress_bar)
        
        self.lbl_progress = QLabel("Prêt")
        self.lbl_progress.setAlignment(Qt.AlignCenter)
        self.lbl_progress.setStyleSheet("color: gray; font-size: 11px;")
        main_layout.addWidget(self.lbl_progress)
        
        main_layout.addWidget(QLabel("Console de progression :"))
        self.console = QTextEdit()
        self.console.setReadOnly(True)
        main_layout.addWidget(self.console)
    
    def choisir_fichier_asm(self):
        chemin, _ = QFileDialog.getOpenFileName(
            self,
            "Sélectionnez l'assemblage principal (.asm)",
            "",
            "Assemblage Solid Edge (*.asm)"
        )
        if chemin:
            self.chemin_asm_edit.setText(chemin)
            self.log(f"Fichier sélectionné : {chemin}", 'info')
            nom_asm = os.path.splitext(os.path.basename(chemin))[0]
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            self.nom_sortie_edit.setText(f"Export_PLM_{nom_asm}_{timestamp}.xlsx")
    
    def on_mode_dft_changed(self, index):
        """Active/désactive le champ dossier DFT selon le mode sélectionné."""
        # Mode 0 = Arborescence uniquement (désactivé)
        # Mode 1 = Dossier spécifique uniquement (activé)
        # Mode 2 = Les deux (activé)
        if index == 1:  # Dossier spécifique uniquement
            self.dossier_dft_edit.setEnabled(True)
            self.btn_parcourir_dft.setEnabled(True)
        elif index == 2:  # Les deux
            self.dossier_dft_edit.setEnabled(True)
            self.btn_parcourir_dft.setEnabled(True)
        else:  # Arborescence uniquement
            self.dossier_dft_edit.setEnabled(False)
            self.btn_parcourir_dft.setEnabled(False)
            self.dossier_dft_edit.clear()

    def choisir_dossier_dft(self):
        dossier = QFileDialog.getExistingDirectory(
            self,
            "Sélectionnez le dossier contenant les plans (.dft)",
            ""
        )
        if dossier:
            self.dossier_dft_edit.setText(dossier)
            self.log(f"Dossier plans sélectionné : {dossier}", 'info')

    def log(self, message, msg_type='info'):
        cursor = self.console.textCursor()
        cursor.movePosition(QTextCursor.End)
        color = QColor('black')
        if msg_type == 'success': color = QColor('green')
        elif msg_type == 'error': color = QColor('red')
        elif msg_type == 'warning': color = QColor('orange')
        self.console.setTextColor(color)
        cursor.insertText(message + '\n')
        self.console.setTextCursor(cursor)
        self.console.ensureCursorVisible()
    
    def lancer_extraction(self):
        if self.extraction_en_cours:
            QMessageBox.warning(self, "Attention", "Une extraction est déjà en cours.")
            return
        
        chemin_asm = self.chemin_asm_edit.text()
        if not chemin_asm:
            QMessageBox.warning(self, "Attention", "Veuillez sélectionner un fichier ASM.")
            return
        
        if not os.path.exists(chemin_asm):
            QMessageBox.critical(self, "Erreur", "Le fichier ASM sélectionné n'existe pas.")
            return
        
        os.makedirs(self.dossier_sortie, exist_ok=True)
        self.extraction_en_cours = True
        self.btn_extraire.setEnabled(False)
        self.btn_extraire.setText("Extraction en cours...")
        
        dossier_dft = self.dossier_dft_edit.text() if self.dossier_dft_edit.text() else None
        
        # Convertir l'index du combo en mode de recherche
        mode_index = self.mode_dft_combo.currentIndex()
        if mode_index == 0:
            mode_recherche = "arborescence"
        elif mode_index == 1:
            mode_recherche = "dossier_specifique"
        else:
            mode_recherche = "les_deux"
        
        self.thread = ExtractionThread(chemin_asm, self.dossier_sortie, self.nom_sortie_edit.text(), dossier_dft, mode_recherche)
        self.thread.log_signal.connect(self.log)
        self.thread.progress_signal.connect(self.update_progress)
        self.thread.finished_signal.connect(self.extraction_terminee)
        self.thread.start()
    
    def update_progress(self, value, maximum, message):
        self.progress_bar.setValue(value)
        self.progress_bar.setMaximum(maximum)
        self.lbl_progress.setText(message)
    
    def extraction_terminee(self):
        self.extraction_en_cours = False
        self.btn_extraire.setEnabled(True)
        self.btn_extraire.setText("Lancer l'extraction")
        self.progress_bar.setValue(0)
        self.lbl_progress.setText("Prêt")

def main():
    app = QApplication(sys.argv)
    window = PLMExtractorGUI()
    window.show()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()